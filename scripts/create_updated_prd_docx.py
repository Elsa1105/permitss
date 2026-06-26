from __future__ import annotations

from pathlib import Path
from typing import Iterable, List, Sequence, Tuple

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT = Path("/Users/joshualim/Documents/Joseph - ePermit web app/Franklin_ePermit_Updated_PRD_v1.1.docx")


BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
LIGHT_BLUE = "E8EEF5"
LIGHT_GRAY = "F2F4F7"
MID_GRAY = "666666"
RISK_RED = "9B1C1C"
OK_GREEN = "1F7A4D"
AMBER = "7A5A00"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_text(cell, text: str, bold: bool = False, color: str | None = None) -> None:
    cell.text = ""
    p = cell.paragraphs[0]
    run = p.add_run(text)
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    for paragraph in cell.paragraphs:
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.line_spacing = 1.05
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_table_widths(table, widths: Sequence[float]) -> None:
    for row in table.rows:
        for idx, width in enumerate(widths):
            if idx < len(row.cells):
                row.cells[idx].width = Inches(width)


def add_table(
    doc: Document,
    headers: Sequence[str],
    rows: Sequence[Sequence[str]],
    widths: Sequence[float] | None = None,
    font_size: float = 9.0,
) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    hdr = table.rows[0]
    set_repeat_table_header(hdr)
    for i, h in enumerate(headers):
        set_cell_text(hdr.cells[i], h, bold=True, color="000000")
        set_cell_shading(hdr.cells[i], LIGHT_GRAY)
    for row_data in rows:
        row = table.add_row()
        for i, value in enumerate(row_data):
            set_cell_text(row.cells[i], value)
            for p in row.cells[i].paragraphs:
                for r in p.runs:
                    r.font.size = Pt(font_size)
    if widths:
        set_table_widths(table, widths)
    doc.add_paragraph()


def add_bullets(doc: Document, items: Iterable[str], level: int = 0) -> None:
    style = "List Bullet" if level == 0 else "List Bullet 2"
    for item in items:
        p = doc.add_paragraph(style=style)
        p.add_run(item)


def add_numbered(doc: Document, items: Iterable[str]) -> None:
    for item in items:
        p = doc.add_paragraph(style="List Number")
        p.add_run(item)


def add_status_run(paragraph, status: str) -> None:
    run = paragraph.add_run(status)
    run.bold = True
    color = {
        "Implemented": OK_GREEN,
        "Partially implemented": AMBER,
        "Not implemented": RISK_RED,
        "Pending decision": AMBER,
    }.get(status, MID_GRAY)
    run.font.color.rgb = RGBColor.from_string(color)


def add_callout(doc: Document, title: str, body: str, fill: str = LIGHT_BLUE) -> None:
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(title)
    r.bold = True
    r.font.color.rgb = RGBColor.from_string(DARK_BLUE)
    p2 = cell.add_paragraph()
    p2.paragraph_format.space_after = Pt(0)
    p2.add_run(body)
    doc.add_paragraph()


def add_use_case(
    doc: Document,
    code: str,
    title: str,
    status: str,
    actors: str,
    goal: str,
    preconditions: Sequence[str],
    main_flow: Sequence[str],
    business_rules: Sequence[str],
    acceptance: Sequence[str],
) -> None:
    doc.add_heading(f"{code}. {title}", level=2)
    p = doc.add_paragraph()
    p.add_run("Status: ").bold = True
    add_status_run(p, status)
    p.add_run(" | Actors: ").bold = True
    p.add_run(actors)
    p = doc.add_paragraph()
    p.add_run("Goal: ").bold = True
    p.add_run(goal)
    if preconditions:
        doc.add_heading("Preconditions", level=3)
        add_bullets(doc, preconditions)
    doc.add_heading("Main Flow", level=3)
    add_numbered(doc, main_flow)
    if business_rules:
        doc.add_heading("Business Rules / Exceptions", level=3)
        add_bullets(doc, business_rules)
    doc.add_heading("Acceptance Criteria", level=3)
    add_bullets(doc, acceptance)


def configure_styles(doc: Document) -> None:
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    for name, size, color, before, after in [
        ("Title", 24, "0B2545", 0, 6),
        ("Subtitle", 12, MID_GRAY, 0, 12),
        ("Heading 1", 16, BLUE, 16, 8),
        ("Heading 2", 13, BLUE, 12, 6),
        ("Heading 3", 12, DARK_BLUE, 8, 4),
    ]:
        style = styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        if name.startswith("Heading"):
            style.font.bold = True


def add_cover(doc: Document) -> None:
    p = doc.add_paragraph()
    p.style = doc.styles["Title"]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.add_run("Franklin Offshore ePermit System")
    p = doc.add_paragraph()
    p.style = doc.styles["Subtitle"]
    p.add_run("Updated Product Requirements Document (PRD) - Hot Work MVP and Client Follow-up Requirements")

    meta = [
        ("Prepared for", "Joseph / Franklin Offshore International Pte Ltd"),
        ("Prepared by", "Codingo Assignments Pte. Ltd."),
        ("Version", "1.1 - Updated Draft"),
        ("Date", "28 May 2026"),
        ("Reference baseline", "Codingo_Franklin_ePermit_PRD.docx, version 1.0 draft dated 30 April 2026"),
        ("Additional inputs", "Meeting audio note and follow-up email thread dated 28 May 2026"),
    ]
    add_table(doc, ["Field", "Value"], meta, widths=[1.7, 4.8], font_size=10)
    add_callout(
        doc,
        "Purpose of This Update",
        "This document consolidates the original Hot Work Permit MVP PRD, the features already implemented in the current demo, implementation gaps identified from code review, and the client follow-up requirements captured in the 28 May 2026 meeting and email thread.",
        fill="EAF2F8",
    )


def add_exec_summary(doc: Document) -> None:
    doc.add_heading("1. Executive Summary", level=1)
    doc.add_paragraph(
        "The current demo already implements the core Hot Work Permit workflow: authenticated access, internal roles, permit creation, Stage I-IV submissions, SRM approval, Day 2-14 endorsements, photo upload and annotation UI, PDF generation, audit log, and admin CSV user import. The follow-up client emails broaden the MVP into a stronger operational workflow with guest contractor access, mandatory risk-assessment documentation, SIMOPS-focused SRM approval, QR-code live permit display, HOD notifications, and company/site-based role control."
    )
    add_callout(
        doc,
        "Updated Direction",
        "Proceed with Hot Work as the first working permit form, but design the platform so it can expand to approximately six permit types. Guest-based contractor access and QR-based live permit display are now part of the target product direction, subject to final workflow and visibility decisions from Franklin.",
    )
    add_table(
        doc,
        ["Theme", "Current Position"],
        [
            ("Phase 1 permit", "Hot Work Permit (Onshore), FOI-SG-057."),
            ("Core workflow", "Applicant -> Assessor -> SRM/Approver -> Close-out, with audit logging."),
            ("Client additions", "Guest applicant, RA/supporting documents, QR live display, email notifications, HOD lists, company/site context."),
            ("Key governance update", "SRM acts as SIMOPS/coordination approver, not physical inspector. No backdating of approvals."),
            ("Major open decision", "Final SRM endorsement workflow and cross-company/site role matrix."),
        ],
        widths=[1.8, 4.7],
        font_size=9.5,
    )


def add_sources_and_status(doc: Document) -> None:
    doc.add_heading("2. Source Inputs and Status Legend", level=1)
    add_bullets(
        doc,
        [
            "Original PRD: Codingo_Franklin_ePermit_PRD.docx, version 1.0 draft.",
            "Current implemented demo in the project folder: Next.js + Supabase app.",
            "Implementation gap review captured in issues.md.",
            "Meeting audio note: Block 437 Woodlands St 41.m4a.",
            "Client follow-up emails: Alex Lim minutes and Joseph Mok response dated 28 May 2026.",
        ],
    )
    add_table(
        doc,
        ["Status", "Meaning"],
        [
            ("Implemented", "Feature exists in the current demo and broadly satisfies the PRD requirement."),
            ("Partially implemented", "Feature exists but is incomplete, under-validated, or missing follow-up requirements."),
            ("Not implemented", "Feature is not present in the current demo."),
            ("Pending decision", "Technical implementation depends on Franklin confirming workflow, roles, content, or policy."),
        ],
        widths=[1.7, 4.8],
        font_size=9.5,
    )


def add_scope(doc: Document) -> None:
    doc.add_heading("3. Updated Scope", level=1)
    doc.add_heading("3.1 Phase 1 Scope", level=2)
    add_bullets(
        doc,
        [
            "Hot Work Permit (Onshore) digital workflow based on FOI-SG-057.",
            "Internal authenticated workflow for applicant, assessor, SRM/approver, and admin.",
            "Guest applicant access for contractors, subject to final role definition.",
            "RA/supporting document upload for contractor and internal submissions.",
            "Photo capture, annotation, and evidence capture.",
            "Printable PDF and QR-based live permit display.",
            "Audit log, admin exports, and email notifications for submission, approval, rejection, overdue, expiry, and escalation events.",
        ],
    )
    doc.add_heading("3.2 Progressive Expansion", level=2)
    add_bullets(
        doc,
        [
            "Approximately six permit types are expected to be implemented progressively after the first Hot Work workflow is stabilised.",
            "Future forms may include hazardous works and other HSE permit categories to be confirmed by Franklin.",
            "The platform data model should support permit_type, company/site context, role-per-site assignments, reusable attachments, QR display, and common audit logging.",
        ],
    )
    doc.add_heading("3.3 Out of Scope Unless Confirmed", level=2)
    add_bullets(
        doc,
        [
            "Direct integration with contractor systems.",
            "Responsibility for contractor internal assessments beyond collecting self-declared and uploaded evidence.",
            "Offline/PWA mode.",
            "HR system sync.",
            "Legal compliance wording beyond client-approved text.",
        ],
    )


def add_feature_matrix(doc: Document) -> None:
    doc.add_heading("4. Feature Status Matrix", level=1)
    rows = [
        ("Authentication and 8-hour session", "Implemented", "Supabase email/password auth and 8-hour JWT config exist."),
        ("Internal roles: Applicant, Assessor, SRM, Admin", "Implemented", "Role enum and core dashboards exist."),
        ("Guest Applicant role", "Not implemented", "Required by client follow-up emails for contractors."),
        ("Company/site context (FOI vs CFE)", "Not implemented", "Emails require company selection and site-scoped roles."),
        ("Hot Work permit creation", "Implemented", "Header fields, validation, and serial generation exist."),
        ("External contractor data fields", "Not implemented", "Need name, company, supervisor registration number, and self-declaration model."),
        ("Stage I checklist", "Implemented", "Three original checklist items enforced."),
        ("RA upload, top controls summary, worker briefing acknowledgment", "Not implemented", "Joseph email requests these for Guest Applicant duty support."),
        ("Photo upload and annotation UI", "Partially implemented", "UI exists; annotation persistence/RLS and PDF rendering need fixes."),
        ("Supporting document upload", "Not implemented", "Need PDF/RA/supporting files beyond photos."),
        ("Stage II assessor fit/not-fit", "Implemented", "Core fit/not-fit and remarks flow exists."),
        ("Assessor condition-verification checklist", "Not implemented", "Need isolations, barricade, gas test, fire watch, PPE, and evidence capture."),
        ("Stage III SRM approval", "Partially implemented", "Approval exists; copy and rules should explicitly say SIMOPS/coordination approval."),
        ("Day 2-14 SRM endorsements", "Partially implemented", "UI and RPC exist; workflow and date controls remain pending."),
        ("Stage IV close-out", "Partially implemented", "Close-out exists but early close-out and override rules need confirmation."),
        ("QR live permit display", "Not implemented", "Client agreed direction; content/security pending."),
        ("Email notifications and HOD lists", "Not implemented", "Submission/approval/rejection plus escalation triggers requested."),
        ("PDF printout", "Partially implemented", "PDF exists; QR and annotated photo rendering missing."),
        ("Audit log", "Partially implemented", "Core audit exists; direct write/auto-expiry gaps need hardening."),
        ("Admin CSV user import", "Implemented", "Bulk import and invite flow exists."),
        ("Admin role/qualification editing", "Partially implemented", "API supports updates; UI mainly toggles active/inactive."),
        ("Permit CSV export", "Implemented", "Admin permit and audit CSV exports exist."),
        ("Additional permit types", "Not implemented", "Progressive expansion to about six permit types expected."),
    ]
    add_table(doc, ["Feature", "Status", "Notes"], rows, widths=[2.0, 1.3, 3.2], font_size=8.5)


def add_roles(doc: Document) -> None:
    doc.add_heading("5. Roles, Companies, and Permissions", level=1)
    add_table(
        doc,
        ["Role", "Primary Responsibilities", "Key Restrictions / Notes", "Status"],
        [
            ("Guest Applicant / Contractor", "Submit permit request, self-declare contractor supervisor details, upload RA/supporting documents, acknowledge worker briefing.", "Cannot act as assessor, SRM, or approver. Validation responsibility for contractor-provided details remains with contractor unless Franklin defines a verification list.", "Not implemented"),
            ("Applicant / Foreman / Supervisor", "Create internal permits, complete Stage I, upload photos/documents, submit, close out where allowed.", "Cannot endorse Stage II or approve Stage III. Cannot sign as another user.", "Implemented"),
            ("Safety Assessor", "Review pending Stage II permits, perform condition verification, record fit/not-fit, add remarks/evidence.", "Should not approve Stage III unless separately authorised. Client email requests condition-verification checklist.", "Partially implemented"),
            ("SRM / Approver", "SIMOPS and coordination approval; ensure no conflict of works; approve/reject/revoke; endorse continuation where workflow requires.", "Not an inspector. Approvals may be remote. No backdating. Site/premises restrictions must prevent cross-site coordination conflicts.", "Partially implemented"),
            ("Admin", "Manage users, roles, qualifications, companies/sites, email lists, exports, audit review, CSV import.", "Cannot sign as operational actor unless explicitly approved by workflow and audit rules.", "Partially implemented"),
        ],
        widths=[1.45, 2.05, 2.25, 0.75],
        font_size=8.3,
    )

    doc.add_heading("5.1 Company and Site Context", level=2)
    add_bullets(
        doc,
        [
            "Users must be associated with one or more company/site contexts, such as FOI and CFE.",
            "Login or session setup should let users select the company/site context when they have more than one assignment.",
            "Roles such as Assessor and SRM must be assigned per company/site, not globally by default.",
            "SRMs remain restricted to specific premises to prevent coordination conflicts across sites.",
        ],
    )


def add_use_cases(doc: Document) -> None:
    doc.add_heading("6. Detailed Feature Use Cases", level=1)
    use_cases = [
        (
            "UC-01",
            "Authenticate User",
            "Implemented",
            "All users",
            "Allow only known users to access the ePermit system under their own identity.",
            ["User has an active Supabase Auth account.", "User is active in the public users table."],
            ["User navigates to login.", "User enters email and password.", "System validates credentials.", "System loads active role, qualifications, and company/site context.", "System redirects user to the dashboard."],
            ["Sessions expire after 8 hours.", "Inactive users are blocked.", "Every operational action must use auth.uid() as the actor."],
            ["Valid active user can log in.", "Inactive user is rejected.", "Unauthenticated requests to protected pages/API routes are blocked."],
        ),
        (
            "UC-02",
            "Manage Users, Roles, and Qualifications",
            "Partially implemented",
            "Admin",
            "Allow Franklin administrators to provision and maintain authorised users.",
            ["Admin is authenticated and active."],
            ["Admin opens Users page.", "Admin imports CSV or edits user data.", "System creates/invites users or updates existing records.", "System stores role, qualification, active status, and future company/site mapping.", "System reflects changes in queues and workflow permissions."],
            ["Guest applicants must be separable from internal applicants.", "Role changes must be audited if required by Franklin.", "Company/site mapping is required before cross-company workflow goes live."],
            ["CSV import works.", "Admin can activate/deactivate users.", "Future UI supports role, qualification, guest, and company/site edits."],
        ),
        (
            "UC-03",
            "Select Company / Site Context",
            "Not implemented",
            "Users with multi-company/site access",
            "Ensure users act only within the correct company/site context, such as FOI or CFE.",
            ["User has more than one company/site assignment."],
            ["User logs in.", "System prompts for company/site context if multiple are available.", "User selects context.", "System filters dashboards, permit visibility, role permissions, and approver lists by selected context."],
            ["SRM permissions must be site-scoped.", "Assessor and SRM assignments must not leak across company/site boundaries.", "Context changes should be explicit and visible in the UI."],
            ["User sees only permits allowed in the selected context.", "SRM cannot approve/coordinate outside assigned premises.", "Audit records include company/site context."],
        ),
        (
            "UC-04",
            "Guest Applicant Submits Permit Request",
            "Not implemented",
            "Guest Applicant / Contractor",
            "Allow external contractors to submit permit requests without using internal employee accounts.",
            ["Guest account exists.", "Contractor has email/password login.", "Permit type is enabled for guest submission."],
            ["Guest logs in.", "System presents contractor permit request form.", "Guest enters work details, name, company, supervisor type, and supervisor registration number.", "Guest uploads RA/supporting documents.", "Guest confirms self-declaration and worker briefing fields.", "System creates draft or submitted permit according to final workflow."],
            ["Guest cannot act as assessor or approver.", "Contractor-provided data is self-declared unless Franklin supplies a validation list.", "Fraudulent entries remain contractor responsibility, but system must preserve audit trail."],
            ["Guest can submit a request.", "Internal assessor/SRM can review it.", "Guest cannot access internal-only queues or approvals."],
        ),
        (
            "UC-05",
            "Create Hot Work Permit Header",
            "Implemented",
            "Applicant, Admin, approved SRM override if confirmed",
            "Capture core FOI-SG-057 permit details and generate a unique serial number.",
            ["User has permission to create Hot Work permit.", "Commencement date is not before today."],
            ["User selects New Hot Work Permit.", "User enters vessel/project, contractor, location, dates, description, and hazards.", "System validates dates and required fields.", "System generates serial number.", "System stores permit in Draft state."],
            ["Current serial format is FOI-HWP-YYYY-NNN but final format remains open.", "Completion date must be on or after commencement date.", "Description minimum is 10 characters."],
            ["Draft permit is created.", "Serial is unique.", "Applicant identity is tied to session."],
        ),
        (
            "UC-06",
            "Submit Stage I Self-Declaration",
            "Partially implemented",
            "Applicant or Guest Applicant",
            "Confirm pre-work safety commitments before assessor review.",
            ["Permit is in Draft state.", "Actor is the applicant of record or permitted override actor."],
            ["User reviews Stage I declaration.", "User confirms ventilation/lighting, permit display with sketch, and watchman/fire extinguisher or hose.", "User uploads required evidence and documents where applicable.", "User submits Stage I.", "System records name, department/company, timestamp, and signature from session."],
            ["All three original checklist items are mandatory.", "For guest applicants, RA upload, top controls summary, and worker briefing acknowledgement should be mandatory.", "SRM/admin override behavior requires final confirmation."],
            ["Stage I cannot submit unless mandatory confirmations are complete.", "Submission moves permit to Pending Safety Assessment.", "Audit records submitted action and actor."],
        ),
        (
            "UC-07",
            "Upload and Annotate Photos / Location Evidence",
            "Partially implemented",
            "Applicant, Guest Applicant, authorised internal users",
            "Capture worksite photos or sketches and mark issues/locations visually.",
            ["Permit is in a state where photo upload/edit is allowed."],
            ["User takes photo or uploads image.", "System stores photo in permit storage path.", "User opens annotation tool.", "User marks arrows, circles, freehand lines, and text labels.", "User saves annotation.", "System shows photo as annotated and includes evidence in permit record."],
            ["Current max is five photos.", "RLS currently needs update permission for annotation_data.", "PDF must render annotation overlays, not only original photos."],
            ["Photo uploads successfully.", "Annotation persists after refresh.", "Annotated evidence appears in permit PDF/export once fixed."],
        ),
        (
            "UC-08",
            "Upload RA and Supporting Documents",
            "Not implemented",
            "Guest Applicant, Applicant, Assessor, Admin",
            "Attach risk assessment, PDFs, and supporting files to the permit record.",
            ["User is authorised for the permit.", "File type and size are allowed."],
            ["User opens Supporting Documents section.", "User uploads RA/PDF/supporting file.", "System scans/validates metadata as configured.", "System stores file with permit-scoped path.", "System displays file name, type, uploader, and timestamp.", "Authorised users download or view the document."],
            ["RA upload should be mandatory for guest contractor submissions unless Franklin says otherwise.", "Documents may be confidential and should not appear in public QR view unless approved.", "Delete rules must align with audit and permit state."],
            ["Documents upload and can be retrieved.", "Unauthorised users cannot access files.", "Audit log records add/remove actions."],
        ),
        (
            "UC-09",
            "Complete Stage II Assessor Verification",
            "Partially implemented",
            "Safety Assessor",
            "Verify conditions and record fit/not-fit outcome before SRM coordination approval.",
            ["Permit is Pending Safety Assessment.", "Assessor has role/qualification for company/site."],
            ["Assessor opens queue.", "Assessor reviews permit details, photos, RA, and supporting documents.", "Assessor completes condition-verification checklist.", "Assessor records evidence where required.", "Assessor marks Fit or Not Fit.", "System requires remarks if Not Fit.", "System moves Fit permits to Pending SRM Approval or Not Fit permits to terminal Not Fit state."],
            ["Joseph email requests checklist items such as isolations, barricade, gas test, fire watch, PPE, and evidence capture.", "Assessor verification should precede activation.", "Assessor remains an inspector/condition verifier, unlike SRM."],
            ["Unqualified user cannot submit Stage II.", "Not Fit requires remarks.", "Fit moves permit to SRM queue.", "Stage II data is tied to authenticated user."],
        ),
        (
            "UC-10",
            "Complete Stage III SIMOPS / Coordination Approval",
            "Partially implemented",
            "SRM / Approver",
            "Approve, reject, or revoke work based on coordination and conflict-of-work controls.",
            ["Permit has passed Stage II as Fit.", "SRM is assigned to the correct company/site."],
            ["SRM reviews permit, RA, assessor checklist, and schedule/location.", "SRM evaluates SIMOPS and no-conflict-of-work conditions.", "SRM approves or rejects.", "System requires reason for rejection.", "Approved permit becomes Approved/Active."],
            ["SRM is a coordination approver, not a physical inspector.", "Approval can be remote.", "SRM cannot backdate approval.", "SRM cannot approve their own raised permit unless Franklin explicitly changes that control."],
            ["SRM approval moves permit to Active.", "Rejection requires reason and is terminal.", "Separation-of-duties guard is enforced according to final policy."],
        ),
        (
            "UC-11",
            "Perform Day 2-14 Endorsement / Continuation",
            "Pending decision",
            "SRM / Approver",
            "Record continuation, rejection, or revocation for multi-day permits according to final Franklin workflow.",
            ["Permit is Active and spans more than one day.", "SRM is authorised for the permit company/site."],
            ["System determines current permit day.", "SRM reviews active permit and coordination status.", "SRM chooses Continue, Reject, or Revoke for the applicable day.", "System requires remarks for Reject/Revoke.", "System records timestamp and actor.", "System updates permit state as required."],
            ["Final workflow is still under discussion.", "No backdating should be allowed.", "Future dating should be blocked unless explicitly allowed.", "Weekend/remote handling needs policy confirmation.", "Approval/endorsement is coordination, not physical inspection."],
            ["System only allows valid day endorsements.", "Duplicate same-day endorsement is blocked.", "Reject/Revoke requires remarks.", "Audit records day, action, actor, and timestamp."],
        ),
        (
            "UC-12",
            "Complete Stage IV Close-out",
            "Partially implemented",
            "Applicant, permitted SRM/Admin override if confirmed",
            "Record notification of completion and archive the permit.",
            ["Permit is eligible for close-out under final date/state rules."],
            ["User opens permit detail.", "System presents close-out confirmation.", "User confirms completion.", "System records closer identity, department/company, date, time, and signature.", "System moves permit to Closed/Completed."],
            ["Current implementation may allow close-out before completion date.", "Override rules need confirmation.", "If applicant unavailable, authorised override and reason should be logged."],
            ["Close-out is final.", "Audit records closed action.", "Permit appears in closed/archive state."],
        ),
        (
            "UC-13",
            "Generate PDF Permit",
            "Partially implemented",
            "All authorised users",
            "Produce a printable permit that resembles FOI-SG-057 and includes the latest permit data.",
            ["User can view the permit."],
            ["User clicks Print PDF.", "System loads permit, stages, endorsements, and photos.", "System generates A4 PDF.", "User opens or prints PDF for worksite display."],
            ["Annotated photo overlays and QR code are not yet rendered.", "PDF layout exactness vs modernised layout remains an open decision.", "Supporting documents may be referenced but not embedded unless required."],
            ["PDF opens from permit detail.", "All completed stages are represented.", "Future version includes QR and annotated evidence."],
        ),
        (
            "UC-14",
            "Scan QR Code for Live Permit Display",
            "Not implemented",
            "Workers, visitors, clients, inspectors, authorised viewers",
            "Allow worksite users to scan a static QR code and view the latest permit status/content.",
            ["Permit has a generated QR URL.", "Franklin has approved the public/private content policy."],
            ["User scans QR code at worksite.", "System opens stable permit URL.", "System displays latest permit status and approved public fields.", "If token/login is required, system enforces access.", "If permit is renewed/replaced, the scanned view reflects the current linked permit or version according to final rule."],
            ["QR may be publicly visible; content must be reviewed for confidentiality.", "Joseph raised phishing/fraudulent QR concerns; anti-tamper design is required.", "Potential mitigations include Franklin domain, QR authenticity instructions, tokenized URLs, and visible serial/status cross-checks."],
            ["QR URL is stable.", "Display updates when permit changes.", "Sensitive attachments/signatures are hidden unless approved.", "Fraudulent QR risk is addressed in design."],
        ),
        (
            "UC-15",
            "Send Email Notifications and Escalations",
            "Not implemented",
            "System, applicants, assessors, SRMs, HODs, admins",
            "Notify relevant parties when action is needed or permit status changes.",
            ["Email provider is configured.", "HOD/recipient lists are provided by Franklin."],
            ["Trigger event occurs, such as submission, approval, rejection, overdue approval, expiry, scope change, incident/near miss, or revocation.", "System determines recipients by role, company/site, permit type, and escalation rules.", "System sends email with relevant links and summary.", "System records send status if required."],
            ["HOD email lists are pre-configured.", "Escalation triggers requested by Joseph include overdue approvals, permit expiry, scope change, incident/near miss.", "Email content should avoid exposing sensitive details if links are public."],
            ["Submission/approval/rejection notifications are sent.", "Overdue/expiry escalations fire according to configured rules.", "Admins can update recipient lists."],
        ),
        (
            "UC-16",
            "View Audit Log and Exports",
            "Partially implemented",
            "Admin, involved permit parties",
            "Maintain accountable history for every significant action.",
            ["User is authorised to view audit data."],
            ["System writes audit events for permit creation, stage submission, fit/not-fit, approval, rejection, revocation, endorsement, close-out, expiry, uploads, deletes, and admin changes.", "Admin opens audit log.", "Admin filters or exports audit CSV."],
            ["Current audit log exists but write_audit must be protected from direct client forgery.", "Auto-expiry and file changes need audit rows.", "Audit should include company/site context once implemented."],
            ["Audit entries are append-only.", "Admin can export CSV.", "Users cannot forge, edit, or delete audit events."],
        ),
        (
            "UC-17",
            "Admin Export Permit and Audit Data",
            "Implemented",
            "Admin",
            "Export permit and audit records for review and reporting.",
            ["Admin is authenticated and active."],
            ["Admin opens audit or permit export endpoint.", "System validates admin role.", "System queries visible records.", "System returns CSV file."],
            ["Export filtering may need expansion for company/site, date range, permit type, and status.", "PDF batch export is not currently implemented."],
            ["Admin can export audit CSV.", "Admin can export permit CSV.", "Non-admin cannot export."],
        ),
        (
            "UC-18",
            "Add Additional Permit Types",
            "Not implemented",
            "Admin, applicant, guest applicant, assessor, SRM",
            "Extend the platform beyond Hot Work to additional HSE permit forms.",
            ["Franklin provides final list and templates.", "Common workflow and type-specific fields are defined."],
            ["Admin enables permit type.", "Applicant selects permit type.", "System displays type-specific header, checklist, documents, and approval workflow.", "Workflow runs through configured stages.", "PDF/QR render the selected permit type."],
            ["About six permit types are expected progressively.", "Some permit types may need different assessor checklists or approval stages.", "Hot Work remains Phase 1."],
            ["New permit type can be configured without duplicating the whole app.", "Permit type appears in dashboards and exports.", "PDF and QR support the selected type."],
        ),
    ]
    for uc in use_cases:
        add_use_case(doc, *uc)


def add_data_model(doc: Document) -> None:
    doc.add_heading("7. Data Model Updates", level=1)
    add_table(
        doc,
        ["Area", "Current", "Required Update"],
        [
            ("users", "id, email, full_name, department, role, qualified_for, active.", "Add guest applicant role and company/site role assignments, or introduce user_company_roles table."),
            ("permits", "Hot Work fields plus applicant/assessor/SRM/closer IDs.", "Add company_id/site_id, external applicant fields, supervisor_registration_number, top_controls_summary, worker briefing fields, QR token/public slug, current_version linkage if needed."),
            ("permit_stages", "Stage I-IV JSON data.", "Extend Stage I and Stage II JSON schemas for RA, controls, worker briefing, condition verification, and evidence."),
            ("permit_photos", "Photos and annotation_data.", "Add update policy, annotation rendering strategy, and audit events."),
            ("attachments", "Not present.", "New table for RA, PDFs, supporting files with file type, storage path, uploader, timestamp, visibility, and state rules."),
            ("notifications", "Not present.", "New tables for recipient groups, event rules, sent email logs, and escalation state."),
            ("QR/live display", "Not present.", "New QR token/public URL fields and public-view policy table or route-level access rules."),
            ("audit_log", "Append-only table exists.", "Harden write path and add event coverage for expiry, file/document changes, admin role changes, QR access events if required."),
            ("permit_types", "Permit type stored as text.", "Consider permit type registry for future forms, templates, workflows, and PDF/QR configuration."),
        ],
        widths=[1.3, 2.35, 2.85],
        font_size=8.5,
    )


def add_security(doc: Document) -> None:
    doc.add_heading("8. Security and Governance Requirements", level=1)
    doc.add_heading("8.1 Identity and Access", level=2)
    add_bullets(
        doc,
        [
            "All operational users, including guest applicants, must authenticate with email/password.",
            "Users can act only under their own logged-in identity.",
            "Role and qualification checks must be enforced server-side and in database functions/RLS, not only in the UI.",
            "Company/site context must restrict permit visibility and approvals.",
        ],
    )
    doc.add_heading("8.2 No Backdating", level=2)
    add_bullets(
        doc,
        [
            "Approvals and endorsements must use system-generated timestamps.",
            "Users must not be able to enter or alter approval dates/times.",
            "Backdating was explicitly identified by the client as fraudulent and must be technically prevented.",
        ],
    )
    doc.add_heading("8.3 QR Security", level=2)
    add_bullets(
        doc,
        [
            "QR codes should point to Franklin/Codingo-approved domains only.",
            "Public QR content must be limited to fields Franklin approves.",
            "The QR view should show enough permit identity to detect tampering, such as serial number, vessel/project, status, validity, and domain guidance.",
            "Sensitive data such as signatures, internal email lists, RA documents, and client-sensitive attachments should not be public unless approved.",
        ],
    )
    doc.add_heading("8.4 Audit Integrity", level=2)
    add_bullets(
        doc,
        [
            "Audit records must be append-only.",
            "Client code must not be able to call a generic audit writer to forge entries.",
            "All state changes, expiry, document/photo changes, notification sends, and admin role changes should be logged where required.",
        ],
    )


def add_non_functional(doc: Document) -> None:
    doc.add_heading("9. Non-Functional Requirements", level=1)
    add_table(
        doc,
        ["Category", "Requirement"],
        [
            ("Performance", "Page load under 2 seconds on 4G where practical; dashboard query under 500 ms p99 for 1,000+ permits; PDF generation under 5 seconds."),
            ("Devices", "Tablet-first, mobile-acceptable, desktop-acceptable. Support latest two iPadOS versions and modern Android tablet browsers."),
            ("Availability", "Online-first MVP. Offline/PWA support remains post-MVP unless Franklin confirms site connectivity requires it."),
            ("Data retention", "Closed/expired permits retained indefinitely by default until Franklin confirms retention/archival policy."),
            ("Storage", "Photos, annotated evidence, RA, PDFs, and supporting documents stored in private object storage with permit-scoped access."),
            ("Usability", "Field workflows should be simple, button/dropdown-led, and avoid unnecessary typing."),
            ("Compliance wording", "WSH/legal framing should be reviewed by Franklin. The system supports duty evidence but is not itself a legal opinion."),
        ],
        widths=[1.45, 5.05],
        font_size=9,
    )


def add_acceptance(doc: Document) -> None:
    doc.add_heading("10. Updated Acceptance Criteria", level=1)
    add_table(
        doc,
        ["ID", "Criterion", "Status"],
        [
            ("AC-01", "Admin can import internal users and assign roles/qualifications.", "Implemented / expand UI"),
            ("AC-02", "Guest applicant can log in and submit contractor permit request with name, company, supervisor registration number.", "Not implemented"),
            ("AC-03", "Guest/internal applicant can upload RA/supporting PDFs and photos.", "Partially implemented"),
            ("AC-04", "Stage I requires original safety checklist plus RA/top-controls/briefing fields where applicable.", "Partially implemented"),
            ("AC-05", "Assessor completes condition-verification checklist and fit/not-fit outcome.", "Partially implemented"),
            ("AC-06", "SRM approval is labelled and recorded as SIMOPS/coordination approval, not inspection.", "Partially implemented"),
            ("AC-07", "No user can backdate approvals or endorsements.", "Partially implemented / needs hardening"),
            ("AC-08", "Day endorsement workflow follows Franklin's final policy and blocks invalid dates.", "Pending decision"),
            ("AC-09", "QR code opens live latest permit display with approved public content only.", "Not implemented"),
            ("AC-10", "Email notifications send for submission, approval, rejection, overdue approval, expiry, scope change, and incident/near miss where configured.", "Not implemented"),
            ("AC-11", "PDF includes latest permit data, stage signatures, day endorsements, QR code, and annotated evidence.", "Partially implemented"),
            ("AC-12", "Audit log cannot be forged and includes all significant actions.", "Partially implemented / needs hardening"),
            ("AC-13", "Company/site context prevents cross-site SRM approval conflicts.", "Not implemented"),
            ("AC-14", "Hot Work prototype is ready for client review before expansion to additional permit types.", "In progress"),
        ],
        widths=[0.65, 4.7, 1.15],
        font_size=8.3,
    )


def add_open_questions(doc: Document) -> None:
    doc.add_heading("11. Open Decisions Required from Franklin", level=1)
    add_table(
        doc,
        ["Decision", "Why It Matters", "Owner / Input Needed"],
        [
            ("Final SRM endorsement workflow", "Determines date validation, Day 2-14 UI, weekend handling, and audit rules.", "Alex, Joseph, Mahadhir, Gilbert, IT"),
            ("SRM override rules", "Current code permits SRM override on Stage I/II/IV, but separation-of-duties controls need final approval.", "HSE / Management"),
            ("QR public content", "Determines what scanned users can see and what must remain login-protected.", "HSE + IT + Alex"),
            ("QR anti-phishing controls", "Joseph raised fraudulent QR risk; mitigations need agreement.", "IT + Vendor"),
            ("Guest applicant field list", "Determines schema and form design.", "HSE"),
            ("Condition-verification checklist", "Determines Stage II form content and evidence capture.", "HSE"),
            ("HOD email lists and escalation rules", "Required for notifications.", "Alex / IT"),
            ("Company/site role mapping", "Required for FOI vs CFE workflow control.", "IT / Admin"),
            ("Permit numbering format", "Current format may differ from paper register.", "HSE / Admin"),
            ("Additional permit type list and templates", "Required for platform expansion beyond Hot Work.", "HSE / Operations"),
            ("Retention and archival policy", "Impacts storage, audit, and data lifecycle.", "Management / HSE"),
        ],
        widths=[1.75, 3.0, 1.75],
        font_size=8.3,
    )


def add_backlog(doc: Document) -> None:
    doc.add_heading("12. Recommended Implementation Backlog", level=1)
    add_table(
        doc,
        ["Priority", "Work Item", "Rationale"],
        [
            ("P0", "Confirm SRM workflow, guest role, company/site matrix, and QR public content.", "These decisions shape schema, RLS, and user experience."),
            ("P0", "Harden audit writer and add missing audit events.", "Audit integrity is core governance."),
            ("P0", "Implement guest applicant role and contractor fields.", "Client direction confirmed in follow-up email."),
            ("P0", "Add RA/supporting document upload.", "Joseph identified RA upload as duty-supporting requirement."),
            ("P1", "Add Stage II condition-verification checklist and evidence capture.", "Closes gap between simple fit/not-fit and WSH duty support."),
            ("P1", "Revise SRM approval wording to SIMOPS/coordination approval.", "Aligns product with client governance model."),
            ("P1", "Implement QR live display with anti-phishing/security controls.", "Client direction confirmed; requires content decision."),
            ("P1", "Implement email notifications, HOD lists, and escalation triggers.", "Client requested notifications and governance escalation."),
            ("P1", "Fix photo annotation persistence and PDF rendering.", "Current annotation UX is not fully end-to-end."),
            ("P2", "Add company/site context and site-scoped roles.", "Needed for FOI/CFE control and cross-site SRM restrictions."),
            ("P2", "Clean up lint/build warnings and Next 16 middleware migration.", "Improves maintainability before wider rollout."),
            ("P3", "Generalise platform for additional permit types.", "Supports expected future scope after Hot Work prototype."),
        ],
        widths=[0.75, 2.9, 2.85],
        font_size=8.5,
    )


def add_appendix(doc: Document) -> None:
    doc.add_heading("Appendix A. Current Demo Implementation Notes", level=1)
    add_bullets(
        doc,
        [
            "Stack: Next.js App Router, TypeScript, Tailwind, Supabase Auth/Postgres/Storage, pdf-lib.",
            "Implemented pages include login, dashboard, permit list/detail, new permit, admin users, and admin audit log.",
            "Implemented APIs include permit creation, Stage I-IV transitions, endorsements, PDF generation, admin CSV import, audit export, and permit export.",
            "Database includes users, permits, permit_stages, permit_endorsements, permit_photos, audit_log, and permit_counters.",
            "Migrations include later SRM authority override changes in 0005_srm_authority.sql.",
            "Verification: npm run typecheck passed; npm run build passed; npm run lint currently fails because next lint is obsolete in this setup.",
        ],
    )
    doc.add_heading("Appendix B. Change Summary from PRD v1.0", level=1)
    add_bullets(
        doc,
        [
            "Guest applicant role added to target scope.",
            "Contractor name, company, and supervisor registration number added to target data capture.",
            "Mandatory RA/supporting document upload added for contractor workflow.",
            "Assessor condition-verification checklist added.",
            "SRM role clarified as SIMOPS/coordination approval rather than inspection.",
            "No-backdating rule elevated to explicit system requirement.",
            "QR live permit display added to target scope.",
            "Email notifications, HOD lists, and escalation triggers added.",
            "Company/site context for FOI vs CFE added.",
            "Additional permit types moved from future-only context to progressive roadmap after Hot Work prototype.",
        ],
    )


def add_footer(doc: Document) -> None:
    section = doc.sections[0]
    footer = section.footer
    p = footer.paragraphs[0]
    p.text = "Franklin ePermit Updated PRD v1.1 - Draft"
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for r in p.runs:
        r.font.size = Pt(8)
        r.font.color.rgb = RGBColor.from_string(MID_GRAY)


def main() -> None:
    doc = Document()
    configure_styles(doc)
    add_footer(doc)
    add_cover(doc)
    doc.add_page_break()
    add_exec_summary(doc)
    add_sources_and_status(doc)
    add_scope(doc)
    add_feature_matrix(doc)
    add_roles(doc)
    add_use_cases(doc)
    add_data_model(doc)
    add_security(doc)
    add_non_functional(doc)
    add_acceptance(doc)
    add_open_questions(doc)
    add_backlog(doc)
    add_appendix(doc)
    doc.core_properties.title = "Franklin Offshore ePermit System - Updated PRD"
    doc.core_properties.subject = "Hot Work MVP, client follow-up requirements, implemented features, missing features, detailed use cases"
    doc.core_properties.author = "Codingo Assignments Pte. Ltd."
    doc.core_properties.comments = "Generated from original PRD, current implementation review, meeting audio note, and client follow-up emails."
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
